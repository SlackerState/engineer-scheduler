"""导入原「工程师排班表.xlsm」，导出为 Excel。"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import logic
from .db import DB, SLOTS, STATUSES


def _d(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def import_workbook(db: DB, data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=False)
    out = {"engineers": 0, "schedule": 0, "dispatches": 0, "orders": 0}
    by_seq = {}

    if "人员" in wb.sheetnames:
        ws = wb["人员"]
        for seq, name, typ, phone, remark in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            name = _s(name)
            if not name:
                continue
            cur = db.one("SELECT id FROM engineers WHERE name=?", (name,))
            eid = db.save_engineer({"id": cur and cur["id"], "name": name, "type": _s(typ) or "外勤",
                                    "phone": _s(phone), "remark": _s(remark)})
            if seq is not None:
                by_seq[int(seq)] = eid
            out["engineers"] += 1
    ordered = [e["id"] for e in db.engineers()]

    if "排班表" in wb.sheetnames:
        ws = wb["排班表"]
        start = ws["B2"].value
        if isinstance(start, datetime):
            start = start.date()
        if isinstance(start, date):
            # 第 6~10 行依次对应「人员」第 1~5 位；从 C 列起每天两列（上午 / 下午）
            for r in range(6, ws.max_row + 1):
                ref = _s(ws.cell(r, 1).value)
                if not ref:
                    continue
                idx = r - 5
                eid = by_seq.get(idx) or (ordered[idx - 1] if idx <= len(ordered) else None)
                if ref and not ref.startswith("="):
                    hit = db.one("SELECT id FROM engineers WHERE name=?", (ref,))
                    eid = hit["id"] if hit else eid
                if not eid:
                    continue
                for c in range(3, ws.max_column + 1):
                    v = _s(ws.cell(r, c).value)
                    if v in STATUSES:
                        day = (start + timedelta(days=(c - 3) // 2)).isoformat()
                        db.set_cell(eid, day, SLOTS[(c - 3) % 2], v)
                        out["schedule"] += 1

    if "派单记录" in wb.sheetnames:
        ws = wb["派单记录"]
        for row in ws.iter_rows(min_row=2, max_col=7, values_only=True):
            day, name, slot, cust, addr, no, hrs = row
            day, name = _d(day), _s(name)
            if not day or not name or _s(slot) not in SLOTS:
                continue
            e = db.one("SELECT id FROM engineers WHERE name=?", (name,))
            if not e:
                continue
            dup = db.one("SELECT id FROM dispatches WHERE day=? AND engineer_id=? AND slot=? AND customer=? "
                         "AND order_no=?", (day, e["id"], _s(slot), _s(cust), _s(no)))
            if dup:
                continue
            hours = float(hrs) if isinstance(hrs, (int, float)) else None
            db.x("INSERT INTO dispatches(day, engineer_id, slot, customer, address, order_no, hours) "
                 "VALUES(?,?,?,?,?,?,?)", (day, e["id"], _s(slot), _s(cust), _s(addr), _s(no), hours))
            out["dispatches"] += 1

    if "订单清单" in wb.sheetnames:
        ws = wb["订单清单"]
        for row in ws.iter_rows(min_row=2, max_col=9, values_only=True):
            day, no, loan, cust, serial, prob, eng, status, prog = row
            vals = [_d(day), _s(no), _s(loan), _s(cust), _s(serial), _s(prob), _s(eng),
                    _s(status) or "待处理", _s(prog)]
            if not any(vals[1:7]):
                continue
            dup = db.one("SELECT id FROM orders WHERE service_order_no=? AND service_order_no<>''", (vals[1],)) \
                or db.one("SELECT id FROM orders WHERE customer=? AND serial=? AND problem=?",
                          (vals[3], vals[4], vals[5]))
            if dup:
                continue
            db.x("INSERT INTO orders(day, service_order_no, loan_no, customer, serial, problem, engineer, "
                 "status, progress) VALUES(?,?,?,?,?,?,?,?,?)", vals)
            out["orders"] += 1
    return out


FILLS = {
    "空闲": "D9F2D9", "在岗": "D9F2D9", "外出": "FFF2B3", "驻点": "D6E6FA",
    "休假": "D9D9D9", "培训": "E6D9F5", "休息": "F0F0F0",
}
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F5597")


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    ws.freeze_panes = "A2"


def export_workbook(db: DB, start: str, days: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "订单清单"
    cols = ["日期", "Udesk 工单号", "服务订单号", "借料单号", "客户", "机器序列号", "问题", "工程师", "状态",
            "处理进度", "D365 状态", "订单类型", "预约上门（考核截止）", "计划上门", "出发", "到达", "现场完工",
            "超期风险", "设备地址", "联系人", "服务网点"]
    _header(ws, cols)
    s = db.settings()
    for o in db.q("SELECT * FROM orders ORDER BY day DESC, id DESC"):
        logic.enrich(o, s)
        plan = f"{o['planned_day']} {o['planned_slot']}".strip()
        ws.append([o["day"], o["udesk_no"], o["service_order_no"], o["loan_no"], o["customer"], o["serial"],
                   o["problem"], o["engineer"], o["status"], o["progress"], o["d365_status"], o["order_type"],
                   o["appointment"], plan, o["setoff_time"], o["arrival_time"], o["complete_time"], o["risk"]["text"],
                   o["address"], o["contact"], o["network"]])
    for i, w in enumerate([11, 12, 12, 10, 22, 18, 50, 16, 8, 30, 10, 8, 18, 14, 16, 16, 16, 14, 36, 24, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("排班表")
    g = logic.schedule_grid(db, start, days)
    ws.append(["姓名", "类型"] + [x for d in g["dates"] for x in (d, "")])
    ws.append(["", ""] + [s for _ in g["dates"] for s in SLOTS])
    for i in range(len(g["dates"])):
        ws.merge_cells(start_row=1, start_column=3 + 2 * i, end_row=1, end_column=4 + 2 * i)
    for r in g["rows"]:
        line = [r["name"], r["type"]]
        for d in g["dates"]:
            for s in SLOTS:
                c = r["cells"][f"{d}|{s}"]
                line.append(logic.effective_status(c["raw"], r["type"], c["n"]))
        ws.append(line)
        for j, cell in enumerate(ws[ws.max_row][2:]):
            c = r["cells"][f"{g['dates'][j // 2]}|{SLOTS[j % 2]}"]
            color = {"dispatched": "FFC983", "conflict": "FF9C9C"}.get(
                c["state"], FILLS.get(logic.base_status(c["raw"], r["type"]), "FFFFFF"))
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center")
    for label, typ in (("外勤可派人数", "外勤"), ("驻点在岗人数", "驻点")):
        ws.append([label, ""] + [g["totals"].get(typ, {}).get(f"{d}|{s}", 0) for d in g["dates"] for s in SLOTS])
    ws.column_dimensions["A"].width = 14


    ws = wb.create_sheet("人员")
    _header(ws, ["序号", "姓名", "类型", "联系电话", "备注", "D365 别名"])
    for i, e in enumerate(db.engineers(), 1):
        ws.append([i, e["name"], e["type"], e["phone"], e["remark"], e["aliases"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_workload(db: DB, day_from: str, day_to: str) -> bytes:
    """工作量统计导出：汇总、工单明细、排班明细三张表。"""
    from . import workload
    data = workload.export_rows(db, day_from, day_to)
    sm = data["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "工作量汇总"
    sched_cols = list(workload.SCHED_KEYS)
    _header(ws, ["工程师", "类型", "上门工单", "上门天数", "完工", "现场工时(h)", "单均工时(h)",
                 f"超 {sm['max_onsite_h']}h 未计", "平均响应(h)", "超期", "周末上门", "配件消耗(件)", "收费金额"]
            + [f"{k}(天)" for k in sched_cols])
    for r in sm["rows"]:
        name = r["name"] + ("（已停用）" if not r["active"] else "") + ("（人员表外）" if not r["in_table"] else "")
        ws.append([name, r["type"], r["visits"], r["visit_days"], r["completed"], r["onsite_h"], r["onsite_avg"],
                   r["onsite_long"], r["response_avg"], r["late"], r["weekend_visits"], r["parts_used"], r["charge"]]
                  + ([r["sched"][k] for k in sched_cols] if r["in_table"] else [""] * len(sched_cols)))
    t = sm["total"]
    ws.append(["合计", "", t["visits"], "", t["completed"], t["onsite_h"], "", t["onsite_long"], "", t["late"],
               t["weekend_visits"], t["parts_used"], t["charge"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"统计区间 {sm['from']} 至 {sm['to']}（{sm['days']} 天）；上门按 实际到达 > 出发 > 计划上门；"
               f"现场工时 = 到达 → 现场完工，单次超过 {sm['max_onsite_h']} 小时的不计入；"
               f"排班只算工作日（{sm['workdays']} 天），按整天算，当天有派单算外出"])
    ws.column_dimensions["A"].width = 26
    for i in range(2, 14 + len(sched_cols)):
        ws.column_dimensions[get_column_letter(i)].width = 11

    ws = wb.create_sheet("工单明细")
    _header(ws, ["工程师", "上门日期", "上门时间来源", "服务订单号", "Udesk 工单号", "客户", "D365 状态",
                 "派工", "到达", "现场完工", "现场工时(h)", "响应(h)", "预约上门（考核截止）", "超期", "配件消耗(件)", "收费金额"])
    for o in data["orders"]:
        v = o["visit"] or {}
        ws.append([o["local"], v.get("day", ""), v.get("source", ""), o["service_order_no"],
                   o["udesk_no"], o["customer"], o["d365_status"], o["dispatch_time"], o["arrival_time"],
                   o["complete_time"], o["onsite_h"], o["response_h"], o["appointment"],
                   o["risk"]["text"] if o["late"] else "", o["parts_used"], o["charge"]])
    for i, w in enumerate([18, 11, 10, 13, 12, 26, 10, 16, 16, 16, 10, 8, 18, 16, 10, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("排班明细")
    _header(ws, ["日期", "星期", "工程师", "排班", "工单数"])
    for line in data["calendar"]:
        ws.append(line)
        c = ws[ws.max_row][3]
        c.fill = PatternFill("solid", fgColor=FILLS.get(c.value, "FFFFFF"))
    ws.column_dimensions["A"].width, ws.column_dimensions["C"].width = 12, 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
